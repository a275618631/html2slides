"""
Test B: Font Strategy (Preserve, Override, Map) and OOXML East Asian (CJK) Typefaces.
"""

import os
import pytest
from pptx import Presentation
from pptx.oxml.ns import qn

from html2slides.config import Config
from html2slides.font_utils import resolve_font, apply_run_font, contains_cjk
from html2slides.browser import BrowserSession
from html2slides.extractor import DOMExtractor
from html2slides.pptx_renderer import PPTXRenderer


def test_font_modes():
    # 1. Preserve Mode
    cfg_preserve = Config()
    cfg_preserve.fonts.mode = "preserve"
    cfg_preserve.fonts.default = "Noto Sans TC"
    font, _ = resolve_font("Arial, sans-serif", cfg_preserve.fonts)
    assert font == "Arial"

    # 2. Override Mode
    cfg_override = Config()
    cfg_override.fonts.mode = "override"
    cfg_override.fonts.default = "Noto Sans TC"
    font, _ = resolve_font("Arial, sans-serif", cfg_override.fonts)
    assert font == "Noto Sans TC"

    # 3. Map Mode
    cfg_map = Config()
    cfg_map.fonts.mode = "map"
    cfg_map.fonts.default = "Noto Sans TC"
    cfg_map.fonts.map = {"Segoe UI": "Noto Sans TC", "Arial": "Aptos"}
    font, _ = resolve_font("Segoe UI, Tahoma", cfg_map.fonts)
    assert font == "Noto Sans TC"
    font2, _ = resolve_font("Arial, Helvetica", cfg_map.fonts)
    assert font2 == "Aptos"


def test_cjk_detection_and_ooxml(tmp_path):
    assert contains_cjk("智慧製造") is True
    assert contains_cjk("Production AI") is False
    assert contains_cjk("MES 系統 24/7") is True

    # Test PPTX OOXML structure
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tx_box = slide.shapes.add_textbox(0, 0, 100, 100)
    p = tx_box.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = "智慧製造與設備可靠度"
    
    apply_run_font(run, "Noto Sans TC", run.text)

    # Check underlying oxml
    rPr = run._r.get_or_add_rPr()
    ea_elem = rPr.find(qn("a:ea"))
    latin_elem = rPr.find(qn("a:latin"))

    assert ea_elem is not None
    assert ea_elem.get("typeface") == "Noto Sans TC"
    assert latin_elem is not None
    assert latin_elem.get("typeface") == "Noto Sans TC"


def test_zh_tw_fixture_rendering(tmp_path):
    fixture_path = os.path.abspath("tests/fixtures/zh_tw.html")
    config = Config()
    config.deck.slide_selector = ".slide"

    with BrowserSession(config) as browser:
        browser.load_html(fixture_path, is_file=True)
        extractor = DOMExtractor(config)
        deck_model, _ = extractor.extract_deck(browser.page)

    out_pptx = str(tmp_path / "test_zh_tw.pptx")
    renderer = PPTXRenderer(deck_model)
    renderer.render(out_pptx)

    prs = Presentation(out_pptx)
    text_corpus = []
    for shape in prs.slides[0].shapes:
        if shape.has_text_frame:
            text_corpus.append(shape.text_frame.text)

    full_text = " ".join(text_corpus)
    assert "智慧製造" in full_text
    assert "設備可靠度" in full_text
    assert "生產系統架構" in full_text
