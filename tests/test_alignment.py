"""
Test A: Paragraph Alignment (Left, Center, Right, Justify) in PPTX.
"""

import os
import pytest
from pptx import Presentation
from pptx.enum.text import PP_ALIGN

from html2slides.config import Config
from html2slides.browser import BrowserSession
from html2slides.extractor import DOMExtractor
from html2slides.pptx_renderer import PPTXRenderer


def test_paragraph_alignment(tmp_path):
    fixture_path = os.path.abspath("tests/fixtures/alignment.html")
    config = Config()
    config.deck.slide_selector = ".slide"

    with BrowserSession(config) as browser:
        browser.load_html(fixture_path, is_file=True)
        extractor = DOMExtractor(config)
        deck_model, warnings = extractor.extract_deck(browser.page)

    assert len(deck_model.slides) == 1
    slide_model = deck_model.slides[0]

    # Render to PPTX
    out_pptx = str(tmp_path / "test_align.pptx")
    renderer = PPTXRenderer(deck_model)
    renderer.render(out_pptx)

    # Read back PPTX and verify alignment
    prs = Presentation(out_pptx)
    slide = prs.slides[0]

    alignments_found = set()
    for shape in slide.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                alignments_found.add(p.alignment)

    assert PP_ALIGN.LEFT in alignments_found, "PP_ALIGN.LEFT not found"
    assert PP_ALIGN.CENTER in alignments_found, "PP_ALIGN.CENTER not found"
    assert PP_ALIGN.RIGHT in alignments_found, "PP_ALIGN.RIGHT not found"
    assert PP_ALIGN.JUSTIFY in alignments_found, "PP_ALIGN.JUSTIFY not found"
