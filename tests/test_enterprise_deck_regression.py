import os
from pptx import Presentation
from html2slides.config import Config
from html2slides.extractor import DOMExtractor
from html2slides.browser import BrowserSession
from html2slides.pptx_renderer import PPTXRenderer
from html2slides.validator import PresentationValidator


def test_enterprise_3slide_regression(tmp_path):
    fixture_path = os.path.abspath("examples/enterprise_ai_deck/deck_sample.html")
    assert os.path.exists(fixture_path), f"Enterprise HTML fixture not found at {fixture_path}"

    config = Config()
    config.deck.slide_selector = ".slide-container"
    config.deck.width_in = 13.333
    config.deck.height_in = 7.5

    # 1. Extraction
    with BrowserSession(config) as browser:
        browser.load_html(fixture_path, is_file=True)
        extractor = DOMExtractor(config)
        deck_model, warnings = extractor.extract_deck(browser.page)

    assert len(deck_model.slides) == 3, f"Expected 3 slides, got {len(deck_model.slides)}"

    # 2. Rendering
    out_pptx = str(tmp_path / "enterprise_regression.pptx")
    renderer = PPTXRenderer(deck_model)
    renderer.render(out_pptx)

    # 3. PPTX Validation
    validator = PresentationValidator(out_pptx, deck_model)
    report = validator.validate(warnings)

    assert report["status"] == "PASS"
    assert report["slides"] == 3
    assert report["shape_elements"] > 5
    assert report["text_elements"] + report["shape_elements"] > 15

    # 4. Content Verification
    prs = Presentation(out_pptx)
    assert len(prs.slides) == 3
    assert abs(prs.slide_width.inches - 13.333) < 0.01
    assert abs(prs.slide_height.inches - 7.5) < 0.01

    all_slide_texts = []
    for s in prs.slides:
        slide_text = " ".join(shape.text_frame.text for shape in s.shapes if shape.has_text_frame)
        all_slide_texts.append(slide_text)

    # Slide 1 Content Check
    assert "Smart Factory" in all_slide_texts[0] or "Architecture" in all_slide_texts[0]
    # Slide 2 Content Check
    assert "Technical Implementation" in all_slide_texts[1] or "Pipeline" in all_slide_texts[1]
    # Slide 3 Content Check
    assert "Traceability" in all_slide_texts[2] or "Quality Assurance" in all_slide_texts[2]
