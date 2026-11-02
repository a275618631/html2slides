"""
Validation and reporting module inspecting generated PPTX and creating conversion_report.json.
"""

import os
import json
from typing import Dict, Any, List, Optional
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN

from .model import DeckModel


class PresentationValidator:
    """Validates generated PPTX structure and creates structured reports."""

    def __init__(self, pptx_path: str, deck_model: Optional[DeckModel] = None):
        self.pptx_path = pptx_path
        self.deck_model = deck_model

    def validate(self, warnings: Optional[List[str]] = None) -> Dict[str, Any]:
        """Inspect PPTX file and collect statistics."""
        if not os.path.exists(self.pptx_path):
            raise FileNotFoundError(f"PPTX file not found: {self.pptx_path}")

        prs = Presentation(self.pptx_path)

        slide_count = len(prs.slides)
        text_elements_count = 0
        shape_elements_count = 0
        image_elements_count = 0
        
        fonts_found = set()
        alignments = {
            "left": 0,
            "center": 0,
            "right": 0,
            "justify": 0
        }

        for slide in prs.slides:
            for shape in slide.shapes:
                # Shape type detection
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    image_elements_count += 1
                elif shape.shape_type == MSO_SHAPE_TYPE.TEXT_BOX:
                    text_elements_count += 1
                elif shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
                    shape_elements_count += 1

                # Text Frame Inspection
                if shape.has_text_frame:
                    tf = shape.text_frame
                    for p in tf.paragraphs:
                        # Alignment
                        if p.alignment == PP_ALIGN.CENTER:
                            alignments["center"] += 1
                        elif p.alignment == PP_ALIGN.RIGHT:
                            alignments["right"] += 1
                        elif p.alignment == PP_ALIGN.JUSTIFY:
                            alignments["justify"] += 1
                        else:
                            if p.text.strip():
                                alignments["left"] += 1

                        for run in p.runs:
                            if run.font.name:
                                fonts_found.add(run.font.name)

        status = "PASS" if slide_count > 0 and (text_elements_count + shape_elements_count > 0) else "FAIL"

        report: Dict[str, Any] = {
            "status": status,
            "slides": slide_count,
            "text_elements": text_elements_count,
            "shape_elements": shape_elements_count,
            "image_elements": image_elements_count,
            "warnings": warnings or [],
            "fonts": sorted(list(fonts_found)),
            "alignments": alignments
        }

        return report

    def save_report(self, report_path: str, report_data: Dict[str, Any]) -> str:
        """Save report to JSON file."""
        os.makedirs(os.path.dirname(os.path.abspath(report_path)), exist_ok=True)
        with open(report_path, "w", encoding="utf-8") as f:
            json.dump(report_data, f, indent=2, ensure_ascii=False)
        return report_path
