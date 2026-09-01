"""
PowerPoint PPTX Renderer translating DeckModel into editable native PowerPoint shapes and textboxes.
"""

import os
import io
import base64
from typing import Optional, List
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

from .model import (
    DeckModel, SlideModel, Element, ShapeElement, TextElement,
    ImageElement, LineElement, Paragraph, TextRun
)
from .font_utils import apply_run_font


class PPTXRenderer:
    """Renders DeckModel into native, fully editable PowerPoint (.pptx) presentation."""

    def __init__(self, deck: DeckModel):
        self.deck = deck
        self.prs = Presentation()
        self.prs.slide_width = Inches(deck.width_in)
        self.prs.slide_height = Inches(deck.height_in)
        self.blank_layout = self.prs.slide_layouts[6]

    def render(self, output_path: str) -> str:
        """Render all slides in DeckModel and save PPTX file."""
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)

        for slide_model in self.deck.slides:
            self._render_slide(slide_model)

        self.prs.save(output_path)
        return output_path

    def _render_slide(self, slide_model: SlideModel):
        """Render a single slide with background and ordered elements."""
        slide = self.prs.slides.add_slide(self.blank_layout)

        # 1. Slide Background
        if slide_model.background_color_rgb:
            bg_shape = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                0, 0,
                Inches(self.deck.width_in),
                Inches(self.deck.height_in)
            )
            bg_shape.fill.solid()
            bg_shape.fill.fore_color.rgb = RGBColor(*slide_model.background_color_rgb)
            bg_shape.line.fill.background()

        # 2. Render sorted elements
        for element in slide_model.elements:
            if isinstance(element, ShapeElement):
                self._render_shape_element(slide, element)
            elif isinstance(element, TextElement):
                self._render_text_element(slide, element)
            elif isinstance(element, ImageElement):
                self._render_image_element(slide, element)
            elif isinstance(element, LineElement):
                self._render_line_element(slide, element)

    def _render_shape_element(self, slide, el: ShapeElement):
        """Render ShapeElement with solid fill, border, optional accent bar, and text."""
        # Choose shape geometry
        if el.shape_type == "rounded_rectangle":
            shape_type = MSO_SHAPE.ROUNDED_RECTANGLE
        elif el.shape_type == "ellipse":
            shape_type = MSO_SHAPE.OVAL
        else:
            shape_type = MSO_SHAPE.RECTANGLE

        shape = slide.shapes.add_shape(
            shape_type,
            Inches(el.x_in),
            Inches(el.y_in),
            Inches(el.width_in),
            Inches(el.height_in)
        )

        # Fill
        if el.fill_color_rgb:
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor(*el.fill_color_rgb)
        else:
            shape.fill.background()

        # Border Line
        if el.border_color_rgb and el.border_width_pt > 0:
            shape.line.color.rgb = RGBColor(*el.border_color_rgb)
            shape.line.width = Pt(max(0.75, el.border_width_pt))
        else:
            shape.line.fill.background()

        # Left Accent Bar (e.g. for takeaway bars, colored card stripes)
        if el.accent_bar_left:
            bar_color = el.accent_bar_left.get("color_rgb", (0, 82, 155))
            bar_w = el.accent_bar_left.get("width_in", 0.05)
            accent_shape = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(el.x_in),
                Inches(el.y_in),
                Inches(bar_w),
                Inches(el.height_in)
            )
            accent_shape.fill.solid()
            accent_shape.fill.fore_color.rgb = RGBColor(*bar_color)
            accent_shape.line.fill.background()

        # Text Frame
        if el.paragraphs:
            self._populate_text_frame(
                shape.text_frame,
                el.paragraphs,
                el.vertical_align,
                el.padding
            )

    def _render_text_element(self, slide, el: TextElement):
        """Render pure TextElement into a native TextBox."""
        tx_box = slide.shapes.add_textbox(
            Inches(el.x_in),
            Inches(el.y_in),
            Inches(el.width_in),
            Inches(el.height_in)
        )
        self._populate_text_frame(
            tx_box.text_frame,
            el.paragraphs,
            el.vertical_align,
            el.padding
        )

    def _populate_text_frame(self, tf, paragraphs: List[Paragraph], vertical_align: str, padding):
        """Fill TextFrame with paragraphs, runs, formatting, and alignments."""
        tf.word_wrap = True

        # Padding / Margins
        tf.margin_top = Inches(max(0.0, padding.top_in))
        tf.margin_right = Inches(max(0.0, padding.right_in))
        tf.margin_bottom = Inches(max(0.0, padding.bottom_in))
        tf.margin_left = Inches(max(0.0, padding.left_in))

        # Vertical Anchor
        if vertical_align == "middle":
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        elif vertical_align == "bottom":
            tf.vertical_anchor = MSO_ANCHOR.BOTTOM
        else:
            tf.vertical_anchor = MSO_ANCHOR.TOP

        for p_idx, p_model in enumerate(paragraphs):
            p = tf.paragraphs[0] if p_idx == 0 else tf.add_paragraph()

            # Alignment
            if p_model.alignment == "center":
                p.alignment = PP_ALIGN.CENTER
            elif p_model.alignment == "right":
                p.alignment = PP_ALIGN.RIGHT
            elif p_model.alignment == "justify":
                p.alignment = PP_ALIGN.JUSTIFY
            else:
                p.alignment = PP_ALIGN.LEFT

            # Spacing
            if p_model.space_before_pt > 0:
                p.space_before = Pt(p_model.space_before_pt)
            if p_model.space_after_pt > 0:
                p.space_after = Pt(p_model.space_after_pt)

            # Runs
            for r_idx, r_model in enumerate(p_model.runs):
                run = p.add_run()
                run.text = r_model.text
                run.font.size = Pt(r_model.font_size_pt)
                run.font.bold = r_model.bold
                run.font.italic = r_model.italic
                run.font.underline = r_model.underline
                run.font.color.rgb = RGBColor(*r_model.color_rgb)

                # OpenXML Typeface configuration (Latin, CJK, Complex Script)
                apply_run_font(run, r_model.font_family, r_model.text)

    def _render_image_element(self, slide, el: ImageElement):
        """Render Picture shape from local path or data URI."""
        img_source = None

        if el.image_bytes:
            img_source = io.BytesIO(el.image_bytes)
        elif el.src.startswith("data:image/"):
            # Base64 data URI
            try:
                header, encoded = el.src.split(",", 1)
                img_data = base64.b64decode(encoded)
                img_source = io.BytesIO(img_data)
            except Exception:
                img_source = None
        elif os.path.exists(el.src):
            img_source = el.src

        if img_source:
            try:
                slide.shapes.add_picture(
                    img_source,
                    Inches(el.x_in),
                    Inches(el.y_in),
                    Inches(el.width_in),
                    Inches(el.height_in)
                )
            except Exception as e:
                print(f"WARN: Failed to add picture {el.src}: {e}")

    def _render_line_element(self, slide, el: LineElement):
        """Render horizontal or vertical line / connector."""
        line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(el.x_in),
            Inches(el.y_in),
            Inches(max(0.01, el.width_in)),
            Inches(max(0.01, el.height_in))
        )
        line.fill.solid()
        line.fill.fore_color.rgb = RGBColor(*el.stroke_color_rgb)
        line.line.fill.background()
