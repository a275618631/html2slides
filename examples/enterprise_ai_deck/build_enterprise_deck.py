#!/usr/bin/env python3
"""
AdvancedTech Applied AI Solution Architect (2026) Presentation Generator
Strict Typography Rule: Minimum Font Size = 14pt, Scaled Proportional Hierarchy
Format: 16:9 Widescreen (13.333" x 7.5")
"""

import os
import sys
import shutil
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# ==============================================================================
# DESIGN SYSTEM & COLOR TOKENS
# ==============================================================================
FONT_HEADING = "Segoe UI"
FONT_BODY = "Segoe UI"

COLOR_BG = RGBColor(248, 250, 252)          # #F8FAFC Canvas Background
COLOR_WHITE = RGBColor(255, 255, 255)       # #FFFFFF Card Background
COLOR_CARD_BORDER = RGBColor(203, 213, 225) # #CBD5E1 Slate Border

COLOR_PRIMARY_BLUE = RGBColor(0, 82, 155)   # #00529B AdvancedTech Blue / Primary
COLOR_BLUE_TINT = RGBColor(238, 244, 251)   # #EEF4FB Blue Tint Fill
COLOR_BLUE_BORDER = RGBColor(185, 213, 243) # #B9D5F3 Light Blue Border

COLOR_TEXT_DARK = RGBColor(30, 41, 59)      # #1E293B Heading / Text Primary
COLOR_TEXT_MUTED = RGBColor(71, 84, 103)    # #475467 Text Secondary / Muted

COLOR_AMBER = RGBColor(194, 65, 12)         # #C2410C Physical Constraint / Control Path
COLOR_AMBER_TINT = RGBColor(255, 247, 237)  # #FFF7ED Amber Tint Fill
COLOR_AMBER_BORDER = RGBColor(253, 186, 116)# #FDBA74 Amber Border

COLOR_GREEN = RGBColor(22, 101, 52)         # #166534 Outcome / Safety
COLOR_GREEN_TINT = RGBColor(236, 253, 243)  # #ECFDF3 Green Tint Fill
COLOR_GREEN_BORDER = RGBColor(187, 247, 208)# #BBF7D0 Green Border

COLOR_SLATE_TINT = RGBColor(241, 245, 249)  # #F1F5F9 Slate Tint Fill


def build_enterprise_deck_min14(output_paths=None):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLOR_BG
        bg.line.fill.background()
        return bg

    def add_header(slide, kicker_text, title_text, page_num):
        # Kicker (14pt)
        kicker_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.28), Inches(10), Inches(0.28))
        tf_k = kicker_box.text_frame
        tf_k.margin_left = tf_k.margin_top = tf_k.margin_right = tf_k.margin_bottom = 0
        p_k = tf_k.paragraphs[0]
        p_k.text = kicker_text.upper()
        p_k.font.name = FONT_HEADING
        p_k.font.size = Pt(14)
        p_k.font.bold = True
        p_k.font.color.rgb = COLOR_TEXT_MUTED

        # Slide Title (28pt)
        title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.56), Inches(10.5), Inches(0.55))
        tf_t = title_box.text_frame
        tf_t.margin_left = tf_t.margin_top = tf_t.margin_right = tf_t.margin_bottom = 0
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.name = FONT_HEADING
        p_t.font.size = Pt(28)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_PRIMARY_BLUE

        # Meta Tag (14pt)
        meta_box = slide.shapes.add_textbox(Inches(10.2), Inches(0.58), Inches(2.533), Inches(0.4))
        tf_m = meta_box.text_frame
        tf_m.margin_left = tf_m.margin_top = tf_m.margin_right = tf_m.margin_bottom = 0
        p_m = tf_m.paragraphs[0]
        p_m.text = f"Che-Yu Wu | {page_num}/6"
        p_m.alignment = PP_ALIGN.RIGHT
        p_m.font.name = FONT_BODY
        p_m.font.size = Pt(14)
        p_m.font.bold = True
        p_m.font.color.rgb = COLOR_TEXT_MUTED

        # Divider Line
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.16), Inches(12.133), Inches(0.02))
        line.fill.solid()
        line.fill.fore_color.rgb = COLOR_CARD_BORDER
        line.line.fill.background()

    def add_takeaway_bar(slide, text):
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(6.45), Inches(12.133), Inches(0.58))
        bar.fill.solid()
        bar.fill.fore_color.rgb = COLOR_SLATE_TINT
        bar.line.color.rgb = COLOR_PRIMARY_BLUE
        bar.line.width = Pt(2.5)
        tf = bar.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = Inches(0.25)
        p = tf.paragraphs[0]
        p.text = text
        p.font.name = FONT_HEADING
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = COLOR_TEXT_DARK

    # =========================================================================
    # SLIDE 1: CAREER CONTINUUM & SCOPE
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)
    add_header(s1, "CAREER CONTINUUM & SCOPE", "From Manufacturing Systems to Production AI Architecture", 1)

    cards_p1 = [
        ("01 FOUNDATION", "MES & Production Systems", "24/7 manufacturing reality, data timing, and production workflow."),
        ("02 INTEGRATION", "Reliability & Automation / PLC", "System integration, failover validation, and production interfaces."),
        ("03 APPLIED AI", "Production AI", "Model predictions meet physical execution constraints."),
        ("04 ARCHITECTURE", "Architecture Thinking", "Explicit boundaries, controlled complexity, and reusable principles.")
    ]

    card_w = Inches(2.85)
    gap = Inches(0.244)
    for i, (num, title, desc) in enumerate(cards_p1):
        x = Inches(0.6) + i * (card_w + gap)
        card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(1.30), card_w, Inches(3.55))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_WHITE
        card.line.color.rgb = COLOR_CARD_BORDER
        card.line.width = Pt(1.5)
        
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_top = Inches(0.20)
        tf.margin_left = tf.margin_right = Inches(0.18)

        p1 = tf.paragraphs[0]
        p1.text = num
        p1.font.name = FONT_HEADING
        p1.font.size = Pt(14)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_PRIMARY_BLUE

        p2 = tf.add_paragraph()
        p2.text = title
        p2.font.name = FONT_HEADING
        p2.font.size = Pt(17)
        p2.font.bold = True
        p2.font.color.rgb = COLOR_TEXT_DARK
        p2.space_before = Pt(8)

        p3 = tf.add_paragraph()
        p3.text = desc
        p3.font.name = FONT_BODY
        p3.font.size = Pt(14)
        p3.font.color.rgb = COLOR_TEXT_MUTED
        p3.space_before = Pt(8)

    why_box = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(4.98), Inches(12.133), Inches(1.34))
    why_box.fill.solid()
    why_box.fill.fore_color.rgb = COLOR_WHITE
    why_box.line.color.rgb = COLOR_CARD_BORDER
    why_box.line.width = Pt(1.5)
    tf_why = why_box.text_frame
    tf_why.word_wrap = True
    tf_why.margin_top = Inches(0.16)
    tf_why.margin_left = Inches(0.25)

    pw1 = tf_why.paragraphs[0]
    pw1.text = "Why this role fits my path"
    pw1.font.name = FONT_HEADING
    pw1.font.size = Pt(15)
    pw1.font.bold = True
    pw1.font.color.rgb = COLOR_PRIMARY_BLUE

    pw2 = tf_why.add_paragraph()
    pw2.text = "Extending my manufacturing-systems foundation into Production AI solution architecture."
    pw2.font.name = FONT_BODY
    pw2.font.size = Pt(16)
    pw2.font.bold = True
    pw2.font.color.rgb = COLOR_TEXT_DARK
    pw2.space_before = Pt(5)

    add_takeaway_bar(s1, "Turning AI into systems that production can trust, operate, and scale.")

    # =========================================================================
    # SLIDE 2: BEFORE AI: MAKE PRODUCTION SYSTEMS RELIABLE
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "PRODUCTION CREDIBILITY & RELIABILITY", "Before AI: Make Production Systems Reliable", 2)

    # Banner (21pt)
    banner = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.28), Inches(12.133), Inches(0.56))
    banner.fill.solid()
    banner.fill.fore_color.rgb = COLOR_BLUE_TINT
    banner.line.color.rgb = COLOR_BLUE_BORDER
    banner.line.width = Pt(1.5)
    tf_b = banner.text_frame
    tf_b.vertical_anchor = MSO_ANCHOR.MIDDLE
    pb = tf_b.paragraphs[0]
    pb.text = "REDUNDANCY  ≠  VALIDATED FAILOVER BEHAVIOR"
    pb.alignment = PP_ALIGN.CENTER
    pb.font.name = FONT_HEADING
    pb.font.size = Pt(21)
    pb.font.bold = True
    pb.font.color.rgb = COLOR_PRIMARY_BLUE

    pipe_steps = ["1. ASSUMPTION", "2. EXPLICIT REQUIREMENT", "3. SWITCHING VALIDATION", "4. PRODUCTION VERIFICATION"]
    for i, step in enumerate(pipe_steps):
        p_box = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6 + i * 3.1), Inches(1.94), Inches(2.85), Inches(0.48))
        p_box.fill.solid()
        p_box.fill.fore_color.rgb = COLOR_WHITE
        p_box.line.color.rgb = COLOR_PRIMARY_BLUE
        p_box.line.width = Pt(1.5)
        tf_p = p_box.text_frame
        tf_p.vertical_anchor = MSO_ANCHOR.MIDDLE
        pp = tf_p.paragraphs[0]
        pp.text = step
        pp.alignment = PP_ALIGN.CENTER
        pp.font.name = FONT_HEADING
        pp.font.size = Pt(14)
        pp.font.bold = True
        pp.font.color.rgb = COLOR_PRIMARY_BLUE

    evidences = [
        ("MES Modernization", "Hidden assumption → explicit switching requirement", "Aligned validation scenarios with Hitachi."),
        ("Maintainability", "Proprietary / EOL dependency → maintainable infra", "Standard servers · TCP/IP · Fiber"),
        ("PSS Rebuild", "Data structure & timing → flexible operator views", "Supports changing reporting and viewing needs.")
    ]
    for i, (title, primary_t, sec_t) in enumerate(evidences):
        c_box = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6 + i * 4.15), Inches(2.54), Inches(3.833), Inches(3.78))
        c_box.fill.solid()
        c_box.fill.fore_color.rgb = COLOR_WHITE
        c_box.line.color.rgb = COLOR_CARD_BORDER
        c_box.line.width = Pt(1.5)
        tf_c = c_box.text_frame
        tf_c.word_wrap = True
        tf_c.margin_top = Inches(0.22)
        tf_c.margin_left = tf_c.margin_right = Inches(0.22)

        p1 = tf_c.paragraphs[0]
        p1.text = title
        p1.font.name = FONT_HEADING
        p1.font.size = Pt(18)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_TEXT_DARK

        p2 = tf_c.add_paragraph()
        p2.text = primary_t
        p2.font.name = FONT_BODY
        p2.font.size = Pt(15)
        p2.font.bold = True
        p2.font.color.rgb = COLOR_PRIMARY_BLUE
        p2.space_before = Pt(10)

        p3 = tf_c.add_paragraph()
        p3.text = sec_t
        p3.font.name = FONT_BODY
        p3.font.size = Pt(14)
        p3.font.color.rgb = COLOR_TEXT_MUTED
        p3.space_before = Pt(8)

    add_takeaway_bar(s2, "Delivered only after: Cold Run → Hot Run → Production Verification")

    # =========================================================================
    # SLIDE 3: I START FROM THE PRODUCTION DECISION — NOT THE MODEL
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "SOLUTION JUDGMENT & SCOPE DISCIPLINE", "I Start from the Production Decision — Not the Model", 3)

    flow_steps = [
        ("1. PRODUCTION DECISION", "What production decision are we trying to improve?"),
        ("2. DETERMINISTIC SOLUTION?", "Can deterministic logic solve it reliably?\n(Rules · MES · Automation / PLC)"),
        ("3. ADDITIONAL AI VALUE?", "Does AI create meaningful additional value?"),
        ("4. LIFECYCLE COMPLEXITY?", "Does that value justify the added model, interface & lifecycle complexity?")
    ]
    for i, (tag, text) in enumerate(flow_steps):
        y = Inches(1.28 + i * 1.26)
        f_box = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), y, Inches(6.8), Inches(1.15))
        f_box.fill.solid()
        f_box.fill.fore_color.rgb = COLOR_WHITE
        f_box.line.color.rgb = COLOR_CARD_BORDER
        f_box.line.width = Pt(1.5)
        tf_f = f_box.text_frame
        tf_f.word_wrap = True
        tf_f.margin_top = Inches(0.12)
        tf_f.margin_left = Inches(0.2)

        p1 = tf_f.paragraphs[0]
        p1.text = tag
        p1.font.name = FONT_HEADING
        p1.font.size = Pt(14)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_PRIMARY_BLUE

        p2 = tf_f.add_paragraph()
        p2.text = text
        p2.font.name = FONT_BODY
        p2.font.size = Pt(14.5)
        p2.font.bold = True
        p2.font.color.rgb = COLOR_TEXT_DARK
        p2.space_before = Pt(3)

    # Right Card 1: Deterministic First (with left accent bar)
    r1 = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7.65), Inches(1.28), Inches(5.083), Inches(2.45))
    r1.fill.solid()
    r1.fill.fore_color.rgb = COLOR_WHITE
    r1.line.color.rgb = COLOR_CARD_BORDER
    r1.line.width = Pt(1.5)
    r1_bar = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7.65), Inches(1.28), Inches(0.08), Inches(2.45))
    r1_bar.fill.solid()
    r1_bar.fill.fore_color.rgb = COLOR_TEXT_MUTED
    r1_bar.line.fill.background()

    tf_r1 = r1.text_frame
    tf_r1.word_wrap = True
    tf_r1.margin_top = Inches(0.3)
    tf_r1.margin_left = Inches(0.35)
    tf_r1.margin_right = Inches(0.25)
    p_r1_t = tf_r1.paragraphs[0]
    p_r1_t.text = "Deterministic First"
    p_r1_t.font.name = FONT_HEADING
    p_r1_t.font.size = Pt(19)
    p_r1_t.font.bold = True
    p_r1_t.font.color.rgb = COLOR_TEXT_DARK
    p_r1_d = tf_r1.add_paragraph()
    p_r1_d.text = "Use simpler, deterministic logic when it reliably solves the production problem."
    p_r1_d.font.name = FONT_BODY
    p_r1_d.font.size = Pt(15)
    p_r1_d.font.color.rgb = COLOR_TEXT_MUTED
    p_r1_d.space_before = Pt(10)

    # Right Card 2: AI with Boundaries (with blue accent bar + blue tint)
    r2 = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7.65), Inches(3.92), Inches(5.083), Inches(2.45))
    r2.fill.solid()
    r2.fill.fore_color.rgb = COLOR_BLUE_TINT
    r2.line.color.rgb = COLOR_BLUE_BORDER
    r2.line.width = Pt(1.5)
    r2_bar = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7.65), Inches(3.92), Inches(0.08), Inches(2.45))
    r2_bar.fill.solid()
    r2_bar.fill.fore_color.rgb = COLOR_PRIMARY_BLUE
    r2_bar.line.fill.background()

    tf_r2 = r2.text_frame
    tf_r2.word_wrap = True
    tf_r2.margin_top = Inches(0.3)
    tf_r2.margin_left = Inches(0.35)
    tf_r2.margin_right = Inches(0.25)
    p_r2_t = tf_r2.paragraphs[0]
    p_r2_t.text = "AI with Boundaries"
    p_r2_t.font.name = FONT_HEADING
    p_r2_t.font.size = Pt(19)
    p_r2_t.font.bold = True
    p_r2_t.font.color.rgb = COLOR_PRIMARY_BLUE
    p_r2_d = tf_r2.add_paragraph()
    p_r2_d.text = "Use AI only when its value justifies the added model, interface, and lifecycle complexity."
    p_r2_d.font.name = FONT_BODY
    p_r2_d.font.size = Pt(15)
    p_r2_d.font.color.rgb = COLOR_TEXT_DARK
    p_r2_d.space_before = Pt(10)

    add_takeaway_bar(s3, "If AI adds more complexity than value, I would not use it.")

    # =========================================================================
    # SLIDE 4: WHEN A GOOD MODEL MEETS PHYSICAL REALITY
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "TECHNICAL CASE 1 — EDGE-TRIMMING OPTIMIZATION", "When a Good Model Meets Physical Reality", 4)

    # Tier 1: Model Layer Container
    t1 = s4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.28), Inches(12.133), Inches(1.45))
    t1.fill.solid()
    t1.fill.fore_color.rgb = COLOR_WHITE
    t1.line.color.rgb = COLOR_PRIMARY_BLUE
    t1.line.width = Pt(1.5)
    
    t1_txt = s4.shapes.add_textbox(Inches(0.8), Inches(1.34), Inches(4.5), Inches(1.3))
    tf1_l = t1_txt.text_frame
    tf1_l.margin_left = tf1_l.margin_top = 0
    p1_k = tf1_l.paragraphs[0]
    p1_k.text = "MODEL LAYER"
    p1_k.font.name = FONT_HEADING
    p1_k.font.size = Pt(14)
    p1_k.font.bold = True
    p1_k.font.color.rgb = COLOR_PRIMARY_BLUE
    p1_sub = tf1_l.add_paragraph()
    p1_sub.text = "Historical data analysis & model comparison"
    p1_sub.font.name = FONT_BODY
    p1_sub.font.size = Pt(14)
    p1_sub.font.color.rgb = COLOR_TEXT_MUTED
    p1_sub.space_before = Pt(4)

    # Metric 1: ~20K Records
    m1_box = s4.shapes.add_textbox(Inches(5.4), Inches(1.32), Inches(2.0), Inches(1.3))
    tf_m1 = m1_box.text_frame
    tf_m1.margin_left = tf_m1.margin_top = 0
    pm1_v = tf_m1.paragraphs[0]
    pm1_v.text = "~20K"
    pm1_v.font.name = FONT_HEADING
    pm1_v.font.size = Pt(34)
    pm1_v.font.bold = True
    pm1_v.font.color.rgb = COLOR_PRIMARY_BLUE
    pm1_l = tf_m1.add_paragraph()
    pm1_l.text = "RECORDS"
    pm1_l.font.name = FONT_HEADING
    pm1_l.font.size = Pt(14)
    pm1_l.font.bold = True
    pm1_l.font.color.rgb = COLOR_TEXT_MUTED

    # Middle Comparison
    cmp_box = s4.shapes.add_textbox(Inches(7.4), Inches(1.54), Inches(2.6), Inches(0.6))
    tf_cmp = cmp_box.text_frame
    tf_cmp.margin_left = tf_cmp.margin_top = 0
    pcmp = tf_cmp.paragraphs[0]
    pcmp.text = "Comparison ➔ XGBoost"
    pcmp.font.name = FONT_HEADING
    pcmp.font.size = Pt(16)
    pcmp.font.bold = True
    pcmp.font.color.rgb = COLOR_PRIMARY_BLUE

    # Metric 2: ~99.8% Correctness
    m2_box = s4.shapes.add_textbox(Inches(10.0), Inches(1.32), Inches(2.6), Inches(1.3))
    tf_m2 = m2_box.text_frame
    tf_m2.margin_left = tf_m2.margin_top = 0
    pm2_v = tf_m2.paragraphs[0]
    pm2_v.text = "~99.8%"
    pm2_v.font.name = FONT_HEADING
    pm2_v.font.size = Pt(34)
    pm2_v.font.bold = True
    pm2_v.font.color.rgb = COLOR_PRIMARY_BLUE
    pm2_l = tf_m2.add_paragraph()
    pm2_l.text = "MODEL PREDICTION CORRECTNESS"
    pm2_l.font.name = FONT_HEADING
    pm2_l.font.size = Pt(14)
    pm2_l.font.bold = True
    pm2_l.font.color.rgb = COLOR_TEXT_MUTED

    # Tier 2: Physical Reality Barrier
    t2 = s4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(2.84), Inches(12.133), Inches(1.28))
    t2.fill.solid()
    t2.fill.fore_color.rgb = COLOR_AMBER_TINT
    t2.line.color.rgb = COLOR_AMBER
    t2.line.width = Pt(2)
    tf2 = t2.text_frame
    tf2.word_wrap = True
    tf2.margin_top = Inches(0.16)
    tf2.margin_left = Inches(0.25)
    p2_k = tf2.paragraphs[0]
    p2_k.text = "⚠️  PHYSICAL REALITY BARRIER — Equipment Resolution ≈ 1 mm"
    p2_k.font.name = FONT_HEADING
    p2_k.font.size = Pt(16.5)
    p2_k.font.bold = True
    p2_k.font.color.rgb = COLOR_AMBER

    p2_m = tf2.add_paragraph()
    p2_m.text = "Material Properties   ·   Order Requirements   ·   Equipment Condition"
    p2_m.font.name = FONT_HEADING
    p2_m.font.size = Pt(15)
    p2_m.font.bold = True
    p2_m.font.color.rgb = COLOR_AMBER
    p2_m.space_before = Pt(6)

    # Tier 3: Execution Layer Container
    t3 = s4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(4.22), Inches(12.133), Inches(2.12))
    t3.fill.solid()
    t3.fill.fore_color.rgb = COLOR_WHITE
    t3.line.color.rgb = COLOR_GREEN
    t3.line.width = Pt(2)
    t3_bar = s4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(4.22), Inches(0.08), Inches(2.12))
    t3_bar.fill.solid()
    t3_bar.fill.fore_color.rgb = COLOR_GREEN
    t3_bar.line.fill.background()

    tf3_h = t3.text_frame
    tf3_h.margin_top = Inches(0.10)
    tf3_h.margin_left = Inches(0.3)
    p3_hk = tf3_h.paragraphs[0]
    p3_hk.text = "EXECUTION LAYER · OEM Baseline + AI Compensation"
    p3_hk.font.name = FONT_HEADING
    p3_hk.font.size = Pt(14.5)
    p3_hk.font.bold = True
    p3_hk.font.color.rgb = COLOR_GREEN
    p3_hsub = tf3_h.add_paragraph()
    p3_hsub.text = "OEM calculation remains; AI adds compensation within production bounds."
    p3_hsub.font.name = FONT_BODY
    p3_hsub.font.size = Pt(14)
    p3_hsub.font.color.rgb = COLOR_TEXT_MUTED
    p3_hsub.space_before = Pt(2)

    # Sub-card 1: Online Range Validation (~93.8%)
    c1 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.85), Inches(4.86), Inches(3.3), Inches(1.36))
    c1.fill.solid()
    c1.fill.fore_color.rgb = COLOR_GREEN_TINT
    c1.line.color.rgb = COLOR_GREEN_BORDER
    c1.line.width = Pt(1.5)
    tf_c1 = c1.text_frame
    tf_c1.margin_top = Inches(0.12)
    tf_c1.margin_left = Inches(0.18)
    pc1_v = tf_c1.paragraphs[0]
    pc1_v.text = "~93.8%"
    pc1_v.font.name = FONT_HEADING
    pc1_v.font.size = Pt(26)
    pc1_v.font.bold = True
    pc1_v.font.color.rgb = COLOR_GREEN
    pc1_l = tf_c1.add_paragraph()
    pc1_l.text = "ONLINE RANGE VALIDATION"
    pc1_l.font.name = FONT_HEADING
    pc1_l.font.size = Pt(14)
    pc1_l.font.bold = True
    pc1_l.font.color.rgb = COLOR_GREEN

    # Sub-card 2: Site Benefit (~NT$2M / yr)
    c2 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.35), Inches(4.86), Inches(3.3), Inches(1.36))
    c2.fill.solid()
    c2.fill.fore_color.rgb = COLOR_GREEN_TINT
    c2.line.color.rgb = COLOR_GREEN_BORDER
    c2.line.width = Pt(1.5)
    tf_c2 = c2.text_frame
    tf_c2.margin_top = Inches(0.12)
    tf_c2.margin_left = Inches(0.18)
    pc2_v = tf_c2.paragraphs[0]
    pc2_v.text = "~NT$2M / yr"
    pc2_v.font.name = FONT_HEADING
    pc2_v.font.size = Pt(26)
    pc2_v.font.bold = True
    pc2_v.font.color.rgb = COLOR_GREEN
    pc2_l = tf_c2.add_paragraph()
    pc2_l.text = "ESTIMATED SITE BENEFIT"
    pc2_l.font.name = FONT_HEADING
    pc2_l.font.size = Pt(14)
    pc2_l.font.bold = True
    pc2_l.font.color.rgb = COLOR_GREEN

    # Sub-card 3: Adoption Flow
    c3 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.85), Inches(4.86), Inches(4.65), Inches(1.36))
    c3.fill.solid()
    c3.fill.fore_color.rgb = COLOR_SLATE_TINT
    c3.line.color.rgb = COLOR_CARD_BORDER
    c3.line.width = Pt(1.5)
    tf_c3 = c3.text_frame
    tf_c3.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf_c3.margin_left = tf_c3.margin_right = Inches(0.15)
    pc3 = tf_c3.paragraphs[0]
    pc3.text = "HMI Advisory ➔ Field Validation ➔ Adoption"
    pc3.alignment = PP_ALIGN.CENTER
    pc3.font.name = FONT_HEADING
    pc3.font.size = Pt(14.5)
    pc3.font.bold = True
    pc3.font.color.rgb = COLOR_PRIMARY_BLUE

    add_takeaway_bar(s4, "A good prediction is not yet a production decision.")

    # =========================================================================
    # SLIDE 5: TURNING AN EXTERNAL AI MODEL INTO A PRODUCTION SOLUTION
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "TECHNICAL CASE 2 — PRODUCTION INTEGRATION ARCHITECTURE", "Turning an External AI Model into a Production Solution", 5)

    # Sub-pill (14.5pt)
    sub_pill = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(1.22), Inches(2.8), Inches(0.38))
    sub_pill.fill.solid()
    sub_pill.fill.fore_color.rgb = COLOR_AMBER_TINT
    sub_pill.line.color.rgb = COLOR_AMBER
    sub_pill.line.width = Pt(1.2)
    tf_sp = sub_pill.text_frame
    tf_sp.vertical_anchor = MSO_ANCHOR.MIDDLE
    psp = tf_sp.paragraphs[0]
    psp.text = "Data Path ≠ Control Path"
    psp.alignment = PP_ALIGN.CENTER
    psp.font.name = FONT_HEADING
    psp.font.size = Pt(14.5)
    psp.font.bold = True
    psp.font.color.rgb = COLOR_AMBER

    # Boundary Container Box
    b_box = s5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.64), Inches(12.133), Inches(4.72))
    b_box.fill.solid()
    b_box.fill.fore_color.rgb = COLOR_WHITE
    b_box.line.color.rgb = COLOR_CARD_BORDER
    b_box.line.width = Pt(1.5)
    b_box.line.dash_style = 2

    lbl_box = s5.shapes.add_textbox(Inches(0.8), Inches(1.68), Inches(5), Inches(0.28))
    tf_lbl = lbl_box.text_frame
    tf_lbl.margin_left = tf_lbl.margin_top = 0
    plbl = tf_lbl.paragraphs[0]
    plbl.text = "PRODUCTION INTEGRATION BOUNDARY"
    plbl.font.name = FONT_HEADING
    plbl.font.size = Pt(14)
    plbl.font.bold = True
    plbl.font.color.rgb = COLOR_TEXT_MUTED

    # Top Node: External R&D Scope
    ext_node = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.6), Inches(1.90), Inches(6.133), Inches(1.15))
    ext_node.fill.solid()
    ext_node.fill.fore_color.rgb = COLOR_SLATE_TINT
    ext_node.line.color.rgb = COLOR_CARD_BORDER
    ext_node.line.width = Pt(1.5)
    tf_ext = ext_node.text_frame
    tf_ext.margin_top = Inches(0.12)
    p_ext1 = tf_ext.paragraphs[0]
    p_ext1.text = "EXTERNAL R&D SCOPE"
    p_ext1.alignment = PP_ALIGN.CENTER
    p_ext1.font.name = FONT_HEADING
    p_ext1.font.size = Pt(14)
    p_ext1.font.bold = True
    p_ext1.font.color.rgb = COLOR_TEXT_MUTED
    p_ext2 = tf_ext.add_paragraph()
    p_ext2.text = "External AI Model"
    p_ext2.alignment = PP_ALIGN.CENTER
    p_ext2.font.name = FONT_HEADING
    p_ext2.font.size = Pt(17)
    p_ext2.font.bold = True
    p_ext2.font.color.rgb = COLOR_TEXT_DARK
    p_ext3 = tf_ext.add_paragraph()
    p_ext3.text = "Model & Algorithm Ownership · Provides Correction"
    p_ext3.alignment = PP_ALIGN.CENTER
    p_ext3.font.name = FONT_BODY
    p_ext3.font.size = Pt(14)
    p_ext3.font.color.rgb = COLOR_TEXT_MUTED

    # Mid Left Node: Data Path
    db_node = s5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(3.18), Inches(4.8), Inches(1.28))
    db_node.fill.solid()
    db_node.fill.fore_color.rgb = COLOR_WHITE
    db_node.line.color.rgb = COLOR_PRIMARY_BLUE
    db_node.line.width = Pt(2.5)
    tf_db = db_node.text_frame
    tf_db.margin_top = Inches(0.12)
    tf_db.margin_left = Inches(0.18)
    pdb1 = tf_db.paragraphs[0]
    pdb1.text = "▲ DATA PATH"
    pdb1.font.name = FONT_HEADING
    pdb1.font.size = Pt(14)
    pdb1.font.bold = True
    pdb1.font.color.rgb = COLOR_PRIMARY_BLUE
    pdb2 = tf_db.add_paragraph()
    pdb2.text = "Designated Database"
    pdb2.font.name = FONT_HEADING
    pdb2.font.size = Pt(16.5)
    pdb2.font.bold = True
    pdb2.font.color.rgb = COLOR_TEXT_DARK
    pdb3 = tf_db.add_paragraph()
    pdb3.text = "▲ Feeds Production Context to External Model"
    pdb3.font.name = FONT_BODY
    pdb3.font.size = Pt(14)
    pdb3.font.bold = True
    pdb3.font.color.rgb = COLOR_PRIMARY_BLUE
    pdb3.space_before = Pt(2)

    db_sub = s5.shapes.add_textbox(Inches(0.9), Inches(4.48), Inches(4.8), Inches(0.28))
    tf_dbs = db_sub.text_frame
    tf_dbs.margin_left = tf_dbs.margin_top = 0
    pdbs = tf_dbs.paragraphs[0]
    pdbs.text = "▲ Actual State & Feedback from Automation"
    pdbs.alignment = PP_ALIGN.CENTER
    pdbs.font.name = FONT_HEADING
    pdbs.font.size = Pt(14)
    pdbs.font.bold = True
    pdbs.font.color.rgb = COLOR_PRIMARY_BLUE

    # Mid Right Node: Control Path
    plc_node = s5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7.633), Inches(3.18), Inches(4.8), Inches(1.28))
    plc_node.fill.solid()
    plc_node.fill.fore_color.rgb = COLOR_WHITE
    plc_node.line.color.rgb = COLOR_AMBER
    plc_node.line.width = Pt(2.5)
    tf_plc = plc_node.text_frame
    tf_plc.margin_top = Inches(0.12)
    tf_plc.margin_left = Inches(0.18)
    pplc1 = tf_plc.paragraphs[0]
    pplc1.text = "▼ CONTROL PATH · AI Correction"
    pplc1.font.name = FONT_HEADING
    pplc1.font.size = Pt(14)
    pplc1.font.bold = True
    pplc1.font.color.rgb = COLOR_AMBER
    pplc2 = tf_plc.add_paragraph()
    pplc2.text = "Automation / PLC"
    pplc2.font.name = FONT_HEADING
    pplc2.font.size = Pt(16.5)
    pplc2.font.bold = True
    pplc2.font.color.rgb = COLOR_TEXT_DARK
    pplc3 = tf_plc.add_paragraph()
    pplc3.text = "Updates Existing 20 Control Setting Values"
    pplc3.font.name = FONT_BODY
    pplc3.font.size = Pt(14)
    pplc3.font.bold = True
    pplc3.font.color.rgb = COLOR_AMBER
    pplc3.space_before = Pt(2)

    plc_sub = s5.shapes.add_textbox(Inches(7.633), Inches(4.48), Inches(4.8), Inches(0.28))
    tf_plcs = plc_sub.text_frame
    tf_plcs.margin_left = tf_plcs.margin_top = 0
    pplcs = tf_plcs.paragraphs[0]
    pplcs.text = "▼ Executes correction through existing control setting values"
    pplcs.alignment = PP_ALIGN.CENTER
    pplcs.font.name = FONT_HEADING
    pplcs.font.size = Pt(14)
    pplcs.font.bold = True
    pplcs.font.color.rgb = COLOR_AMBER

    # Bottom Node: Production Equipment
    eq_node = s5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(4.80), Inches(11.533), Inches(0.65))
    eq_node.fill.solid()
    eq_node.fill.fore_color.rgb = COLOR_SLATE_TINT
    eq_node.line.color.rgb = COLOR_CARD_BORDER
    eq_node.line.width = Pt(1.5)
    tf_eq = eq_node.text_frame
    tf_eq.vertical_anchor = MSO_ANCHOR.MIDDLE
    peq = tf_eq.paragraphs[0]
    peq.text = "Production Equipment"
    peq.alignment = PP_ALIGN.CENTER
    peq.font.name = FONT_HEADING
    peq.font.size = Pt(16)
    peq.font.bold = True
    peq.font.color.rgb = COLOR_TEXT_DARK

    # Guardrail / Annotation Box
    annot = s5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(5.55), Inches(11.533), Inches(0.72))
    annot.fill.solid()
    annot.fill.fore_color.rgb = COLOR_WHITE
    annot.line.color.rgb = COLOR_GREEN
    annot.line.width = Pt(1.5)
    tf_an = annot.text_frame
    tf_an.margin_top = Inches(0.1)
    tf_an.margin_left = Inches(0.2)
    pan1 = tf_an.paragraphs[0]
    pan1.text = "• Direct Control Path: Avoids an unnecessary control loop through MES."
    pan1.font.name = FONT_BODY
    pan1.font.size = Pt(14)
    pan1.font.bold = True
    pan1.font.color.rgb = COLOR_PRIMARY_BLUE
    pan2 = tf_an.add_paragraph()
    pan2.text = "• Failure Boundary: PLC timeout detection; OEM baseline and existing control protections remain available."
    pan2.font.name = FONT_BODY
    pan2.font.size = Pt(14)
    pan2.font.bold = True
    pan2.font.color.rgb = COLOR_GREEN
    pan2.space_before = Pt(3)

    add_takeaway_bar(s5, "AI can lose optimization — but it must never take away basic production capability.")

    # =========================================================================
    # SLIDE 6: FROM PROJECTS TO ARCHITECTURE THINKING
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "SYSTEM BOUNDARIES & REUSABLE METHODOLOGY", "From Projects to Architecture Thinking", 6)

    # Left Container
    b_left = s6.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.30), Inches(7.0), Inches(5.05))
    b_left.fill.solid()
    b_left.fill.fore_color.rgb = COLOR_WHITE
    b_left.line.color.rgb = COLOR_CARD_BORDER
    b_left.line.width = Pt(1.5)

    # Sub-card 1: Camera Network
    c_cam = s6.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.82), Inches(1.45), Inches(6.56), Inches(0.90))
    c_cam.fill.solid()
    c_cam.fill.fore_color.rgb = COLOR_WHITE
    c_cam.line.color.rgb = COLOR_CARD_BORDER
    c_cam.line.width = Pt(1.5)
    c_cam.line.dash_style = 2
    tf_cam = c_cam.text_frame
    tf_cam.margin_top = Inches(0.1)
    tf_cam.margin_left = Inches(0.15)
    pcam1 = tf_cam.paragraphs[0]
    pcam1.text = "CAMERA-SIDE NETWORK (~100 CAMERAS)"
    pcam1.font.name = FONT_HEADING
    pcam1.font.size = Pt(14)
    pcam1.font.bold = True
    pcam1.font.color.rgb = COLOR_PRIMARY_BLUE
    pcam2 = tf_cam.add_paragraph()
    pcam2.text = "Camera Server / Storage (Integration Endpoint)"
    pcam2.font.name = FONT_HEADING
    pcam2.font.size = Pt(15.5)
    pcam2.font.bold = True
    pcam2.font.color.rgb = COLOR_TEXT_DARK
    pcam2.space_before = Pt(3)

    # Moxa Firewall Connector Box
    fw_box = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.82), Inches(2.45), Inches(6.56), Inches(0.48))
    fw_box.fill.solid()
    fw_box.fill.fore_color.rgb = COLOR_SLATE_TINT
    fw_box.line.color.rgb = COLOR_CARD_BORDER
    fw_box.line.width = Pt(1.2)
    tf_fw = fw_box.text_frame
    tf_fw.vertical_anchor = MSO_ANCHOR.MIDDLE
    pfw = tf_fw.paragraphs[0]
    pfw.text = "▼  ──[ Moxa Industrial Firewall / NAT · Restricts Raw Video ]──  ▼"
    pfw.alignment = PP_ALIGN.CENTER
    pfw.font.name = FONT_HEADING
    pfw.font.size = Pt(14)
    pfw.font.bold = True
    pfw.font.color.rgb = COLOR_TEXT_MUTED

    # Sub-card 2: MES / PSS Network
    c_mes = s6.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.82), Inches(3.03), Inches(6.56), Inches(0.88))
    c_mes.fill.solid()
    c_mes.fill.fore_color.rgb = COLOR_WHITE
    c_mes.line.color.rgb = COLOR_CARD_BORDER
    c_mes.line.width = Pt(1.5)
    tf_mes = c_mes.text_frame
    tf_mes.margin_top = Inches(0.1)
    tf_mes.margin_left = Inches(0.15)
    pmes1 = tf_mes.paragraphs[0]
    pmes1.text = "MES / PSS NETWORK"
    pmes1.font.name = FONT_HEADING
    pmes1.font.size = Pt(14)
    pmes1.font.bold = True
    pmes1.font.color.rgb = COLOR_PRIMARY_BLUE
    pmes2 = tf_mes.add_paragraph()
    pmes2.text = "Communicates with the Camera Server endpoint only."
    pmes2.font.name = FONT_HEADING
    pmes2.font.size = Pt(15)
    pmes2.font.bold = True
    pmes2.font.color.rgb = COLOR_TEXT_DARK
    pmes2.space_before = Pt(3)

    # Sub-card 3: Context Index
    c_idx = s6.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.82), Inches(4.01), Inches(6.56), Inches(0.92))
    c_idx.fill.solid()
    c_idx.fill.fore_color.rgb = COLOR_BLUE_TINT
    c_idx.line.color.rgb = COLOR_BLUE_BORDER
    c_idx.line.width = Pt(1.5)
    tf_idx = c_idx.text_frame
    tf_idx.margin_top = Inches(0.1)
    tf_idx.margin_left = Inches(0.15)
    pidx1 = tf_idx.paragraphs[0]
    pidx1.text = "PRODUCTION CONTEXT INDEX"
    pidx1.font.name = FONT_HEADING
    pidx1.font.size = Pt(14)
    pidx1.font.bold = True
    pidx1.font.color.rgb = COLOR_PRIMARY_BLUE
    pidx2 = tf_idx.add_paragraph()
    pidx2.text = "Coil ID   ·   Timestamp   ·   Area ID   ·   Camera ID"
    pidx2.font.name = FONT_HEADING
    pidx2.font.size = Pt(16)
    pidx2.font.bold = True
    pidx2.font.color.rgb = COLOR_PRIMARY_BLUE
    pidx2.space_before = Pt(3)

    # Status Chips
    chip1 = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.82), Inches(5.03), Inches(3.15), Inches(0.48))
    chip1.fill.solid()
    chip1.fill.fore_color.rgb = COLOR_GREEN_TINT
    chip1.line.color.rgb = COLOR_GREEN_BORDER
    chip1.line.width = Pt(1.5)
    tf_cp1 = chip1.text_frame
    tf_cp1.vertical_anchor = MSO_ANCHOR.MIDDLE
    pcp1 = tf_cp1.paragraphs[0]
    pcp1.text = "LIVE: Retrieval & Context Indexing"
    pcp1.alignment = PP_ALIGN.CENTER
    pcp1.font.name = FONT_HEADING
    pcp1.font.size = Pt(14)
    pcp1.font.bold = True
    pcp1.font.color.rgb = COLOR_GREEN

    chip2 = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.18), Inches(5.03), Inches(3.20), Inches(0.48))
    chip2.fill.solid()
    chip2.fill.fore_color.rgb = COLOR_SLATE_TINT
    chip2.line.color.rgb = COLOR_CARD_BORDER
    chip2.line.width = Pt(1.5)
    tf_cp2 = chip2.text_frame
    tf_cp2.vertical_anchor = MSO_ANCHOR.MIDDLE
    pcp2 = tf_cp2.paragraphs[0]
    pcp2.text = "NOT DEPLOYED: AI Recognition Model"
    pcp2.alignment = PP_ALIGN.CENTER
    pcp2.font.name = FONT_HEADING
    pcp2.font.size = Pt(14)
    pcp2.font.bold = True
    pcp2.font.color.rgb = COLOR_TEXT_MUTED

    # Right Container: 3 Principles (with blue left accent bar)
    b_right = s6.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7.8), Inches(1.30), Inches(4.933), Inches(5.05))
    b_right.fill.solid()
    b_right.fill.fore_color.rgb = COLOR_WHITE
    b_right.line.color.rgb = COLOR_CARD_BORDER
    b_right.line.width = Pt(1.5)
    b_right_bar = s6.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7.8), Inches(1.30), Inches(0.08), Inches(5.05))
    b_right_bar.fill.solid()
    b_right_bar.fill.fore_color.rgb = COLOR_PRIMARY_BLUE
    b_right_bar.line.fill.background()

    tf_br = b_right.text_frame
    tf_br.word_wrap = True
    tf_br.margin_top = Inches(0.2)
    tf_br.margin_left = Inches(0.32)

    pbr_t = tf_br.paragraphs[0]
    pbr_t.text = "PRINCIPLES I CARRY FORWARD"
    pbr_t.font.name = FONT_HEADING
    pbr_t.font.size = Pt(15)
    pbr_t.font.bold = True
    pbr_t.font.color.rgb = COLOR_PRIMARY_BLUE

    principles = [
        ("1. Context", "Connect data to production context."),
        ("2. Boundary", "Keep external systems useful, but bounded."),
        ("3. Reuse", "Reuse architecture contracts — not site implementations.")
    ]
    for p_title, p_desc in principles:
        pr1 = tf_br.add_paragraph()
        pr1.text = p_title
        pr1.font.name = FONT_HEADING
        pr1.font.size = Pt(18)
        pr1.font.bold = True
        pr1.font.color.rgb = COLOR_TEXT_DARK
        pr1.space_before = Pt(16)

        pr2 = tf_br.add_paragraph()
        pr2.text = p_desc
        pr2.font.name = FONT_BODY
        pr2.font.size = Pt(15)
        pr2.font.color.rgb = COLOR_TEXT_MUTED
        pr2.space_before = Pt(3)

    add_takeaway_bar(s6, "Turning AI from a demonstration into a solution that production can trust, maintain, and scale — step by step.")

    # Save to all destination paths
    if output_paths:
        for path in output_paths:
            os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
            prs.save(path)
            print(f"[OK] Saved presentation to: {os.path.abspath(path)}")
    return prs


if __name__ == "__main__":
    paths = [
        "output/enterprise_deck.pptx",
        "output/enterprise_deck.pptx",
        "output/enterprise_deck.pptx",
        "output/enterprise_deck.pptx",
        "output/enterprise_deck.pptx",
        "output/enterprise_deck.pptx"
    ]
    build_enterprise_deck_min14(paths)
